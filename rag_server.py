from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument

import shutil
import os

# =========================
# FASTAPI
# =========================

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================
# FOLDERS
# =========================

UPLOAD_DIR = "uploads"
VECTOR_DB_DIR = "vector_db"

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(VECTOR_DB_DIR, exist_ok=True)

# =========================
# EMBEDDING MODEL
# =========================

embedding_model = HuggingFaceEmbeddings(
    model_name="sentence-transformers/all-MiniLM-L6-v2"
)

# =========================
# VECTOR DATABASE
# =========================

vectorstore = Chroma(
    persist_directory=VECTOR_DB_DIR,
    embedding_function=embedding_model
)

# =========================
# TEXT SPLITTER
# =========================

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=800,
    chunk_overlap=100
)

# =========================
# EXTRACT TEXT
# =========================

def extract_text(filepath):

    ext = os.path.splitext(filepath)[1].lower()

    text = ""

    # PDF
    if ext == ".pdf":

        reader = PdfReader(filepath)

        for page in reader.pages:

            extracted = page.extract_text()

            if extracted:
                text += extracted + "\n"

    # DOCX
    elif ext == ".docx":

        doc = DocxDocument(filepath)

        for para in doc.paragraphs:
            text += para.text + "\n"

    # TXT
    elif ext == ".txt":

        with open(filepath, "r", encoding="utf-8") as f:
            text = f.read()

    return text

# =========================
# UPLOAD DOCUMENT
# =========================

@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):

    filepath = os.path.join(
        UPLOAD_DIR,
        file.filename
    )

    # SAVE FILE

    with open(filepath, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    # EXTRACT TEXT

    extracted_text = extract_text(filepath)

    if not extracted_text.strip():

        return {
            "success": False,
            "message": "No text extracted"
        }

    # SPLIT INTO CHUNKS

    chunks = text_splitter.split_text(extracted_text)

    docs = []

    for chunk in chunks:

        docs.append(
            Document(
                page_content=chunk,
                metadata={
                    "source": file.filename
                }
            )
        )

    # STORE IN VECTOR DB

    vectorstore.add_documents(docs)
    vectorstore.persist()

    return {
        "success": True,
        "filename": file.filename,
        "chunks": len(chunks)
    }

# =========================
# SEARCH
# =========================

@app.get("/search")

async def search(query:str):

    try:

        docs = vectorstore.similarity_search(
            query,
            k=3
        )

        results = []

        for doc in docs:

            results.append({

                "text":
                doc.page_content,

                "metadata":
                doc.metadata

            })

        return {
            "success": True,
            "results": results
        }

    except Exception as e:

        return {
            "success": False,
            "results": [],
            "error": str(e)
        }

# =========================
# LIST DOCUMENTS
# =========================

@app.get("/documents")
def documents():

    files = os.listdir(UPLOAD_DIR)

    return {
        "documents": files
    }

# =========================
# DELETE DOCUMENT
# =========================

@app.delete("/delete/{filename}")
def delete_document(filename: str):

    filepath = os.path.join(
        UPLOAD_DIR,
        filename
    )

    if os.path.exists(filepath):
        os.remove(filepath)

    return {
        "success": True
    }

# =========================
# ROOT
# =========================

@app.get("/")
def home():

    return {
        "message": "Sarvam AI RAG Server Running"
    }